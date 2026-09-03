"""XLSX ingestion via openpyxl (MIT).

Two things here are deliberate and easy to get wrong:

1. **`data_only=True`, and never evaluating a formula ourselves.** This reads
   the value Excel itself last cached for a formula cell. openpyxl does not
   compute formulas, and neither do we -- that would be Python inventing a
   fact the document doesn't literally contain, the same rule #1 violation as
   the LLM computing a percentage. A workbook saved by a tool that never
   cached results returns None for those cells, so they simply aren't emitted;
   the model then reports the field as null with a reason, which is the honest
   answer rather than a number we made up.

2. **Cells are addressed, not laid out.** `document_text` emits one
   `[B5] value` line per populated cell under a `[ชีท: name]` header, so the
   model can cite `Source.kind="cell_xlsx"` with `ref="Sheet1!B5"` -- A1
   notation, which is exact, unambiguous, and what a human types into Excel's
   name box to go check it. There is no attempt to reconstruct a visual grid
   in the text: a spreadsheet's meaning lives in its cell addresses, and
   faking a table layout in prose would only invite the model to misread
   which value belongs to which label.
"""
from __future__ import annotations

import io
from datetime import date, datetime, time

import openpyxl

from app.ingest.base import IngestedDocument, IngestMeta
from app.models.schema import SheetCell, SheetPreview, SourcePreview
from app.thai.normalize import normalize_text


class XlsxTooLargeError(RuntimeError):
    """Raised instead of silently truncating. Dropping cells past a cap would
    make a value look absent when it is merely unread -- rule #5's "never
    return a partial result quietly" applied to ingestion."""


# A pathological workbook could otherwise blow the whole token budget on cells
# nobody asked about. Chosen to comfortably cover real procurement attachments
# (typically hundreds of cells) while refusing a 100k-cell data dump.
MAX_CELLS = 5_000
# The preview only feeds the viewer, so it can be capped per sheet
# independently -- and unlike the text, truncating it is safe as long as the
# UI says so (SheetPreview.truncated).
MAX_PREVIEW_CELLS_PER_SHEET = 2_000


def _cell_to_text(value: object) -> str | None:
    """Stringify a cell for the prompt. Formatting only -- never arithmetic."""
    if value is None:
        return None
    if isinstance(value, bool):
        # Checked before the numeric branch: bool is a subclass of int.
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        # Excel stores most numbers as floats; "7672800" reads better than
        # "7672800.0" and is the same number, so drop a meaningless .0 only.
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    text = normalize_text(str(value)).strip()
    return text or None


def ingest_xlsx_bytes(data: bytes) -> IngestedDocument:
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    try:
        return _ingest_workbook(workbook)
    finally:
        workbook.close()


def _ingest_workbook(workbook: "openpyxl.Workbook") -> IngestedDocument:
    lines: list[str] = []
    sheets: list[SheetPreview] = []
    total_cells = 0

    for worksheet in workbook.worksheets:
        lines.append(f"\n[ชีท: {worksheet.title}]")
        preview_cells: list[SheetCell] = []
        sheet_cell_count = 0

        for row in worksheet.iter_rows():
            for cell in row:
                text = _cell_to_text(cell.value)
                if text is None:
                    continue
                total_cells += 1
                if total_cells > MAX_CELLS:
                    raise XlsxTooLargeError(
                        f"ไฟล์ Excel นี้มีเซลล์ที่มีข้อมูลมากกว่า {MAX_CELLS:,} เซลล์ "
                        f"ซึ่งเกินขอบเขตที่เครื่องมือนี้อ่านได้ในครั้งเดียว — "
                        f"กรุณาแยกเฉพาะชีท/ช่วงที่เกี่ยวข้องกับ TOR แล้วอัปโหลดใหม่"
                    )
                lines.append(f"[{cell.coordinate}] {text}")
                sheet_cell_count += 1
                if len(preview_cells) < MAX_PREVIEW_CELLS_PER_SHEET:
                    preview_cells.append(SheetCell(ref=cell.coordinate, value=text))

        sheets.append(
            SheetPreview(
                name=worksheet.title,
                cells=preview_cells,
                truncated=sheet_cell_count > len(preview_cells),
            )
        )

    return IngestedDocument(
        document_text="\n".join(lines),
        meta=IngestMeta(
            document_kind="xlsx",
            sheet_names=[ws.title for ws in workbook.worksheets],
        ),
        preview=SourcePreview(document_kind="xlsx", sheets=sheets),
    )
