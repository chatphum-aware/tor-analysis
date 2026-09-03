"""Format detection and the PDF ingestion facade.

Detection is by magic bytes, never filename -- these tests pin that, because
trusting the extension would send a renamed DOCX to the PDF parser and produce
a confusing "couldn't open PDF" instead of "wrong format".
"""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest

from app.ingest.base import UnsupportedDocumentError, UnsupportedKindError
from app.ingest.detect import sniff_document_kind
from app.ingest.pdf_ingest import ingest_meta_from_scan_report, ingest_pdf_path
from app.pdf.extract import build_document_text, extract_document
from app.pdf.scanned import DocumentScanReport, PageReport

SAMPLES_DIR = Path(__file__).resolve().parents[2] / "outputs" / "samples"


def _ooxml_bytes(*entry_names: str) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name in entry_names:
            zf.writestr(name, "<x/>")
    return buf.getvalue()


def test_sniffs_pdf_by_magic_bytes():
    assert sniff_document_kind(b"%PDF-1.7\nrest of file") == "pdf"


def test_sniffs_docx_and_xlsx_by_ooxml_entry_names():
    assert sniff_document_kind(_ooxml_bytes("word/document.xml")) == "docx"
    assert sniff_document_kind(_ooxml_bytes("xl/workbook.xml")) == "xlsx"


def test_filename_never_overrides_the_bytes():
    """A DOCX renamed to .pdf is still a DOCX, and a PDF named .xlsx is still
    a PDF -- the filename is only ever used in error messages."""
    assert sniff_document_kind(_ooxml_bytes("word/document.xml"), "invoice.pdf") == "docx"
    assert sniff_document_kind(b"%PDF-1.4 x", "report.xlsx") == "pdf"


@pytest.mark.parametrize(
    "data",
    [
        pytest.param(b"this is not a document at all", id="plain_text"),
        pytest.param(b"", id="empty"),
        pytest.param(b"PK\x03\x04truncated-not-a-real-zip", id="zip_magic_but_corrupt"),
        pytest.param(_ooxml_bytes("some/other/file.txt"), id="valid_zip_but_not_office"),
    ],
)
def test_rejects_unrecognized_bytes(data):
    with pytest.raises(UnsupportedDocumentError):
        sniff_document_kind(data, "whatever.pdf")


def test_unsupported_error_carries_filename_and_head_for_a_useful_message():
    with pytest.raises(UnsupportedDocumentError) as caught:
        sniff_document_kind(b"nonsense bytes here", "mystery.bin")
    assert caught.value.filename == "mystery.bin"
    assert caught.value.head == b"nonsense bytes he"[:16]


def test_ingest_meta_marks_pdf_and_lists_image_only_page_numbers():
    report = DocumentScanReport(
        page_count=3,
        pages=[
            PageReport(page_number=1, status="text", text_char_count=500, image_coverage_ratio=0.0),
            PageReport(page_number=2, status="image_only", text_char_count=0, image_coverage_ratio=0.9),
            PageReport(page_number=3, status="mixed", text_char_count=200, image_coverage_ratio=0.4),
        ],
    )
    meta = ingest_meta_from_scan_report(report)

    assert meta.document_kind == "pdf"
    assert meta.page_count == 3
    # image_only_pages is a list of page NUMBERS here, unlike the scan report's
    # same-named property, which is a count.
    assert meta.image_only_pages == [2]
    assert meta.usable_text_page_ratio == pytest.approx(2 / 3, abs=1e-4)
    # Fields belonging to other formats stay null rather than being faked.
    assert meta.paragraph_count is None
    assert meta.sheet_names is None


@pytest.mark.skipif(not SAMPLES_DIR.exists(), reason="sample PDFs not present")
def test_facade_document_text_is_byte_identical_to_the_pre_refactor_path():
    """The ingest layer must not change a single byte of what the model sees.
    The `[หน้า N]` markers are what let it cite a page number, so any drift
    here silently corrupts every citation."""
    samples = sorted(SAMPLES_DIR.glob("*.pdf"))
    if not samples:
        pytest.skip("no sample PDFs available")
    for pdf in samples:
        direct = build_document_text(extract_document(str(pdf)).blocks)
        assert ingest_pdf_path(str(pdf)).document_text == direct, f"drift on {pdf.name}"


# --- XLSX ingestion ---------------------------------------------------------
#
# A spreadsheet has no pages, so its citations are sheet+cell in A1 notation.
# The rule #1 case worth guarding: openpyxl with data_only=True returns None
# for a formula cell whose result the saving tool never cached, and we must
# leave it out rather than evaluate the formula ourselves.

import openpyxl  # noqa: E402  (grouped with the XLSX tests it belongs to)

from app.ingest.dispatch import ingest_bytes, supported_kinds  # noqa: E402
from app.ingest.xlsx_ingest import MAX_CELLS, XlsxTooLargeError, ingest_xlsx_bytes  # noqa: E402


def _workbook_bytes(build) -> bytes:
    wb = openpyxl.Workbook()
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_emits_addressed_cells_under_sheet_headers():
    def build(wb):
        ws = wb.active
        ws.title = "งบประมาณ"
        ws["B3"] = "วงเงินงบประมาณ"
        ws["C3"] = 7672800.00

    ing = ingest_xlsx_bytes(_workbook_bytes(build))

    assert "[ชีท: งบประมาณ]" in ing.document_text
    assert "[B3] วงเงินงบประมาณ" in ing.document_text
    # A meaningless trailing .0 is dropped -- formatting only, same number.
    assert "[C3] 7672800" in ing.document_text
    assert ing.meta.document_kind == "xlsx"
    assert ing.meta.sheet_names == ["งบประมาณ"]


def test_xlsx_reports_no_page_metadata_rather_than_faking_it():
    ing = ingest_xlsx_bytes(_workbook_bytes(lambda wb: wb.active.__setitem__("A1", "x")))
    assert ing.meta.page_count is None
    assert ing.meta.usable_text_page_ratio is None
    assert ing.meta.image_only_pages is None


def test_xlsx_omits_formula_cell_with_no_cached_value():
    """rule #1: openpyxl returns None for an uncached formula result, and we
    must not evaluate it. The label cell stays, so the model sees the field
    exists and reports it null with a reason -- rather than us inventing the
    number Excel would have shown."""
    def build(wb):
        ws = wb.active
        ws["B7"] = "ห้าเปอร์เซ็นต์ของวงเงิน"
        ws["C7"] = "=C3*0.05"

    ing = ingest_xlsx_bytes(_workbook_bytes(build))

    assert "[B7] ห้าเปอร์เซ็นต์ของวงเงิน" in ing.document_text
    assert "[C7]" not in ing.document_text
    assert "0.05" not in ing.document_text  # the formula text itself never leaks


def test_xlsx_preview_carries_cells_for_the_viewer_to_highlight():
    def build(wb):
        ws = wb.active
        ws.title = "S1"
        ws["A1"] = "label"
        ws["B1"] = 42
        wb.create_sheet("S2")["A1"] = "other"

    ing = ingest_xlsx_bytes(_workbook_bytes(build))

    assert ing.preview is not None
    assert ing.preview.document_kind == "xlsx"
    assert [s.name for s in ing.preview.sheets] == ["S1", "S2"]
    first = {c.ref: c.value for c in ing.preview.sheets[0].cells}
    assert first == {"A1": "label", "B1": "42"}
    assert all(not s.truncated for s in ing.preview.sheets)


def test_xlsx_refuses_a_workbook_larger_than_the_cell_cap():
    """Fails loudly rather than truncating: a dropped cell would make a value
    look absent when it was merely unread (rule #5's spirit)."""
    def build(wb):
        ws = wb.active
        for i in range(MAX_CELLS + 10):
            ws.cell(row=i + 1, column=1, value=f"v{i}")

    with pytest.raises(XlsxTooLargeError):
        ingest_xlsx_bytes(_workbook_bytes(build))


def test_xlsx_is_disabled_by_default(monkeypatch):
    """XLSX ingestion works (tests above) but is gated off: no real Thai gov
    TOR was found in .xlsx, and Excel attachments (BOQ/price sheets) don't map
    onto this tool's TOR schema. Kept behind a flag rather than deleted --
    see app/ingest/dispatch.py."""
    monkeypatch.delenv("TOR_ENABLE_XLSX", raising=False)
    assert "xlsx" not in supported_kinds()
    data = _workbook_bytes(lambda wb: wb.active.__setitem__("A1", "hello"))
    with pytest.raises(UnsupportedKindError) as caught:
        ingest_bytes(data, "book.xlsx")
    # Names the format it recognized, so the message can say "not supported
    # yet" rather than the misleading "unrecognized file".
    assert caught.value.kind == "xlsx"


def test_dispatch_routes_a_real_xlsx_when_the_flag_is_on(monkeypatch):
    monkeypatch.setenv("TOR_ENABLE_XLSX", "1")
    assert "xlsx" in supported_kinds()
    data = _workbook_bytes(lambda wb: wb.active.__setitem__("A1", "hello"))
    # Deliberately a misleading filename -- the bytes decide.
    ing = ingest_bytes(data, "actually_a_spreadsheet.pdf")
    assert ing.meta.document_kind == "xlsx"


# --- DOCX ingestion ---------------------------------------------------------
#
# Shaped by a real Thai อบต. TOR (134 raw paragraphs + 1 table). The two things
# most worth guarding, both verified against that document:
#   - table cell text must be captured (python-docx's document.paragraphs
#     omits it entirely, and that sample's deliverables live in its table);
#   - body ORDER must be preserved across paragraphs and tables.

from docx import Document as _DocxDocument  # noqa: E402

from app.ingest.docx_ingest import MAX_BLOCKS, DocxTooLargeError, ingest_docx_bytes  # noqa: E402


def _docx_bytes(build) -> bytes:
    doc = _DocxDocument()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_emits_paragraph_refs_the_model_can_copy_verbatim():
    def build(doc):
        doc.add_paragraph("1. ความเป็นมา")
        doc.add_paragraph("3.3 มีวุฒิการศึกษาไม่ต่ำกว่าระดับ ปวช.")

    ing = ingest_docx_bytes(_docx_bytes(build))

    assert "[para:1] 1. ความเป็นมา" in ing.document_text
    assert "[para:2] 3.3 มีวุฒิการศึกษาไม่ต่ำกว่าระดับ ปวช." in ing.document_text
    assert ing.meta.document_kind == "docx"
    assert ing.meta.paragraph_count == 2
    # No page concept in Word -- not faked.
    assert ing.meta.page_count is None
    assert ing.meta.usable_text_page_ratio is None


def test_docx_captures_table_cells_that_document_paragraphs_would_miss():
    """The regression this guards: python-docx's `document.paragraphs`
    excludes table cell content, so a walker built on it would silently drop
    the deliverables table from a real TOR."""
    def build(doc):
        doc.add_paragraph("ก่อนตาราง")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "ลักษณะงาน"
        table.cell(0, 1).text = "ตัวชี้วัดปริมาณงาน"
        table.cell(1, 0).text = "งานบันทึกข้อมูล"
        table.cell(1, 1).text = "ไม่น้อยกว่า 30 รายการ"
        doc.add_paragraph("หลังตาราง")

    ing = ingest_docx_bytes(_docx_bytes(build))

    assert "[t1:r1:c1] ลักษณะงาน" in ing.document_text
    assert "[t1:r2:c2] ไม่น้อยกว่า 30 รายการ" in ing.document_text
    kinds = [b.kind for b in ing.preview.blocks]
    assert kinds.count("table_cell") == 4


def test_docx_preserves_document_order_across_paragraphs_and_tables():
    def build(doc):
        doc.add_paragraph("first")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "in-table"
        doc.add_paragraph("last")

    refs = [b.ref for b in ingest_docx_bytes(_docx_bytes(build)).preview.blocks]

    assert refs == ["para:1", "t1:r1:c1", "para:2"]


def test_docx_normalizes_thai_digits_including_mixed_runs():
    """The real sample writes "3๐ รายการ" -- Arabic 3 beside Thai ๐, meaning
    30 -- and numbers sections "๔.๑.๑". Reusing the PDF path's normalize_text
    handles both."""
    ing = ingest_docx_bytes(_docx_bytes(lambda d: d.add_paragraph("๔.๑.๑ ไม่น้อยกว่า 3๐ รายการ")))
    assert "[para:1] 4.1.1 ไม่น้อยกว่า 30 รายการ" in ing.document_text


def test_docx_skips_blank_paragraphs_so_refs_stay_dense():
    def build(doc):
        doc.add_paragraph("real one")
        doc.add_paragraph("   ")
        doc.add_paragraph("")
        doc.add_paragraph("real two")

    ing = ingest_docx_bytes(_docx_bytes(build))

    assert [b.ref for b in ing.preview.blocks] == ["para:1", "para:2"]
    assert ing.meta.paragraph_count == 2


def test_docx_preview_refs_match_the_refs_in_the_prompt_text():
    """The viewer highlights by matching Source.ref against preview block
    refs, so a drift between the two would break click-to-verify silently."""
    def build(doc):
        doc.add_paragraph("alpha")
        doc.add_table(rows=1, cols=2).rows[0].cells[1].text = "beta"

    ing = ingest_docx_bytes(_docx_bytes(build))

    for block in ing.preview.blocks:
        assert f"[{block.ref}] " in ing.document_text


def test_docx_refuses_a_document_larger_than_the_block_cap():
    def build(doc):
        for i in range(MAX_BLOCKS + 5):
            doc.add_paragraph(f"ย่อหน้า {i}")

    with pytest.raises(DocxTooLargeError):
        ingest_docx_bytes(_docx_bytes(build))


def test_docx_is_enabled_by_default_unlike_xlsx():
    """Real Thai gov TORs do circulate as .docx (verified against a live
    sample), which is why this one is not behind a flag."""
    assert "docx" in supported_kinds()
    data = _docx_bytes(lambda d: d.add_paragraph("hello"))
    assert ingest_bytes(data, "tor.docx").meta.document_kind == "docx"


# --- Vision rendering for scanned PDF pages (Phase 5) -----------------------
#
# The real finding that shaped this: a real 16-page scanned document (under
# MAX_VISION_PAGES=20) still produced ~55MB of base64 PNG data at 200 DPI --
# Anthropic rejected the request outright with a 413 before even checking the
# API key. JPEG cut the same document to ~5MB. Page COUNT alone never bounded
# this; only a byte-size check does, which is why there are two independent
# caps below.

import fitz  # noqa: E402  (PyMuPDF, grouped with the vision tests it builds fixtures for)

import app.ingest.pdf_ingest as pdf_ingest_module  # noqa: E402
from app.ingest.pdf_ingest import TooManyVisionPagesError, ingest_pdf_bytes  # noqa: E402


def _image_only_pdf_bytes(n_pages: int = 1) -> bytes:
    """A synthetic PDF with no text layer at all, classified image_only by
    the same heuristic real scanned pages are -- see app/pdf/scanned.py."""
    doc = fitz.open()
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 800, 1000))
    pix.set_rect(pix.irect, (200, 200, 200))
    for _ in range(n_pages):
        page = doc.new_page()
        page.insert_image(page.rect, pixmap=pix)
    data = doc.tobytes()
    doc.close()
    return data


def test_scanned_pdf_renders_pages_as_jpeg_not_png():
    """PNG was the original choice and is wrong for this content -- see the
    module comment on VISION_RENDER_DPI for why (scanned pages are
    photo-like, which is exactly what PNG's lossless compression handles
    worst)."""
    ing = ingest_pdf_bytes(_image_only_pdf_bytes(1))

    assert len(ing.images) == 1
    assert ing.images[0].media_type == "image/jpeg"
    assert ing.images[0].data[:3] == b"\xff\xd8\xff"  # JPEG magic bytes
    assert ing.images[0].page == 1


def test_text_only_pdf_never_triggers_rendering():
    """The common case (a real text layer, no scanned pages) must be
    completely unaffected -- rendering only runs for image_only pages."""
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "hello, this page has real text")
    data = doc.tobytes()
    doc.close()

    ing = ingest_pdf_bytes(data)

    assert ing.images == []


def test_scanned_pdf_refuses_past_the_page_count_cap(monkeypatch):
    monkeypatch.setattr(pdf_ingest_module, "MAX_VISION_PAGES", 2)

    with pytest.raises(TooManyVisionPagesError):
        ingest_pdf_bytes(_image_only_pdf_bytes(3))


def test_scanned_pdf_refuses_past_the_payload_size_cap_even_under_page_cap(monkeypatch):
    """The regression this guards: a document can be well under
    MAX_VISION_PAGES and still be too large in bytes -- confirmed live (see
    module comment). A cap on page count alone would have missed it."""
    monkeypatch.setattr(pdf_ingest_module, "MAX_VISION_PAGES", 10)
    monkeypatch.setattr(pdf_ingest_module, "MAX_VISION_PAYLOAD_BYTES", 1000)

    with pytest.raises(TooManyVisionPagesError):
        ingest_pdf_bytes(_image_only_pdf_bytes(2))  # 2 pages, well under the page cap


def test_scanned_pdf_within_both_caps_succeeds(monkeypatch):
    monkeypatch.setattr(pdf_ingest_module, "MAX_VISION_PAGES", 5)
    monkeypatch.setattr(pdf_ingest_module, "MAX_VISION_PAYLOAD_BYTES", 50 * 1024 * 1024)

    ing = ingest_pdf_bytes(_image_only_pdf_bytes(3))

    assert len(ing.images) == 3
    assert [img.page for img in ing.images] == [1, 2, 3]
